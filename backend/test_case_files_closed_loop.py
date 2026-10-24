"""用例文件在线编辑 - 闭环测试（不依赖外部 LLM）
运行: cd backend && python test_case_files_closed_loop.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
import pandas as pd
from fastapi.testclient import TestClient

from app.database import init_db

init_db()  # TestClient 不触发 lifespan，需显式建表
from app.utils.seed import create_admin_user

create_admin_user()
from app.main import app
from app.services import captcha_service, case_file_service

captcha_service.verify_captcha = lambda cid, code: True
client = TestClient(app)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  -> {detail}" if not cond and detail else ""))


def login():
    r = client.post("/api/v1/auth/login", json={
        "username": "admin", "password": "admin123", "captcha_id": "x", "captcha_code": "x"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['data']['access_token']}"}


def make_xlsx(rows=3, cols=4, text="用例"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "功能用例"
    for r in range(rows):
        for c in range(cols):
            ws.cell(row=r + 1, column=c + 1, value=f"{text}-{r}-{c}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    admin = login()
    r = client.post("/api/v1/projects", headers=admin, json={
        "name": "_用例文件闭环测试_临时", "product_line": "smart_home", "description": "可删"})
    pid = r.json()["data"]["id"]
    try:
        # 1-2 上传 xlsx / csv
        r = client.post("/api/v1/case-files/upload", headers=admin,
                        files={"file": ("嵌入式用例.xlsx", make_xlsx(), XLSX_MIME)},
                        data={"project_id": str(pid)})
        check("上传 xlsx", r.status_code == 200, r.text[:150])
        fid = r.json()["data"]["id"]

        r = client.post("/api/v1/case-files/upload", headers=admin,
                        files={"file": ("回归用例.csv", "用例编号,标题\nTC-001,登录\n".encode("utf-8"), "text/csv")},
                        data={"project_id": str(pid)})
        check("上传 csv(自动转xlsx)", r.status_code == 200, r.text[:150])
        fid_csv = r.json()["data"]["id"]

        # 3 非法格式
        r = client.post("/api/v1/case-files/upload", headers=admin,
                        files={"file": ("说明.txt", b"hello", "text/plain")},
                        data={"project_id": str(pid)})
        check("非法格式被拒", r.status_code == 400)

        # 4 列表与锁状态
        r = client.get("/api/v1/case-files", headers=admin, params={"project_id": pid})
        items = r.json()["data"]
        check("列表 2 个文件且未锁定", r.status_code == 200 and len(items) == 2
              and all(not i["lock"]["locked"] for i in items))

        # 5-6 抢锁 / 他人只读
        r = client.post(f"/api/v1/case-files/{fid}/open", headers=admin)
        check("打开抢锁成功", r.json()["data"]["lock"]["acquired"])
        other = case_file_service.try_acquire_lock(fid, 999, "李四")
        check("他人抢锁失败(只读)", not other["acquired"] and other["holder_name"] == "admin")

        # 7-8 心跳 / 持锁保存
        check("心跳续锁", client.post(f"/api/v1/case-files/{fid}/heartbeat", headers=admin).status_code == 200)
        new_bytes = make_xlsx(5, 5, "改")
        r = client.post(f"/api/v1/case-files/{fid}/save", headers=admin,
                        files={"file": ("嵌入式用例.xlsx", new_bytes, XLSX_MIME)})
        check("持锁保存", r.status_code == 200, r.text[:150])

        # 9 下载验证落盘
        r = client.get(f"/api/v1/case-files/{fid}/download", headers=admin)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        changed = any(str(c.value or "").startswith("改") for row in wb.active.iter_rows() for c in row)
        check("下载且保存内容落盘", r.status_code == 200 and changed)

        # 10-11 锁失效拒绝保存/心跳
        case_file_service._LOCKS[fid]["expires_at"] = 0
        r = client.post(f"/api/v1/case-files/{fid}/save", headers=admin,
                        files={"file": ("嵌入式用例.xlsx", new_bytes, XLSX_MIME)})
        check("锁失效保存被拒(409)", r.status_code == 409, str(r.status_code))
        check("锁失效心跳被拒(409)", client.post(f"/api/v1/case-files/{fid}/heartbeat", headers=admin).status_code == 409)

        # 12 重新抢锁 + 释放
        r = client.post(f"/api/v1/case-files/{fid}/open", headers=admin)
        check("重新抢锁", r.json()["data"]["lock"]["acquired"])
        check("主动释放", client.post(f"/api/v1/case-files/{fid}/release", headers=admin).status_code == 200)

        # 13 csv 转换内容
        r = client.get(f"/api/v1/case-files/{fid_csv}/download", headers=admin)
        df = pd.read_excel(io.BytesIO(r.content), header=None)
        check("csv 内容正确转换", df.iloc[1, 0] == "TC-001" and df.iloc[1, 1] == "登录")

        # 14 删除
        check("admin 删除", client.delete(f"/api/v1/case-files/{fid_csv}", headers=admin).status_code == 200)
        check("删除后 404", client.get(f"/api/v1/case-files/{fid_csv}/download", headers=admin).status_code == 404)
    finally:
        # 清理：项目删除联动清 case_file
        r = client.delete(f"/api/v1/projects/{pid}", headers=admin)
        check("项目删除联动清理用例文件", r.status_code == 200, r.text[:150])
        print(f"\n结果: {len(PASS)} 通过 / {len(FAIL)} 失败")
        sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
