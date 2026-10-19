"""知识库对话内容导出：Markdown 回答 → Word / PPT / PDF

- Word : python-docx 逐块渲染（标题/列表/表格/段落/代码块，含 **粗体** 内联解析）
- PDF  : reportlab（复用需求体检 PDF 的中文字体方案）
- PPT  : 两段式——LLM 先把回答压缩为幻灯大纲 JSON，python-pptx 按平台深蓝模板渲染
         （内容与排版分离：AI 负责"讲什么"，模板负责"长什么样"）
"""
import io
import json
import logging
import re

logger = logging.getLogger(__name__)

# ================= 轻量 Markdown 块解析（三种格式共用） =================

_TABLE_RE = re.compile(r"^\s*\|.*\|\s*$")
_H_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_LIST_RE = re.compile(r"^\s*(?:[-*•]|\d+[.、)])\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_bold(text: str) -> list[tuple[str, bool]]:
    parts, pos = [], 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts or [(text, False)]


def parse_md_blocks(md: str) -> list[dict]:
    """把 Markdown 解析为块列表（够用即可，不追求完整语法）"""
    blocks: list[dict] = []
    lines = (md or "").replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "text": "\n".join(code)})
            continue
        if _TABLE_RE.match(line):
            rows = []
            while i < n and _TABLE_RE.match(lines[i]):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue
        m = _H_RE.match(s)
        if m:
            blocks.append({"type": f"h{min(len(m.group(1)), 3)}", "text": m.group(2).strip()})
            i += 1
            continue
        if s.startswith(">"):
            blocks.append({"type": "para", "text": "「" + s.lstrip("> ").strip() + "」"})
            i += 1
            continue
        m = _LIST_RE.match(line)
        if m:
            blocks.append({"type": "bullet", "text": m.group(1).strip()})
            i += 1
            continue
        blocks.append({"type": "para", "text": s})
        i += 1
    return blocks


def _ensure_font_reportlab():
    from app.services.report_pdf import _FONT, _ensure_font
    _ensure_font()
    return _FONT


def _md_to_pdf(md: str, title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _ensure_font_reportlab()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm)
    navy = colors.HexColor("#1E2761")
    h1 = ParagraphStyle("h1", fontName=font, fontSize=18, textColor=navy, spaceBefore=10, spaceAfter=5)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=14, textColor=navy, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=16, textColor=colors.HexColor("#1E293B"))
    bullet_style = ParagraphStyle("b", parent=body, leftIndent=12)

    def esc(s):
        return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def rich(text: str) -> str:
        out = []
        for frag, bold in _strip_bold(text):
            frag = esc(frag)
            out.append(f"<b>{frag}</b>" if bold else frag)
        return "".join(out)

    elems: list = []
    elems.append(Paragraph(esc(title), ParagraphStyle("t", fontName=font, fontSize=20, textColor=navy, spaceAfter=8)))
    elems.append(Paragraph(f"<font size='9' color='#64748B'>来源：知识库对话 · {esc(title)}</font>", body))
    elems.append(Spacer(1, 4 * mm))

    for b in parse_md_blocks(md):
        if b["type"] == "h1":
            elems.append(Paragraph(rich(b["text"]), h1))
        elif b["type"] == "h2":
            elems.append(Paragraph(rich(b["text"]), h2))
        elif b["type"] == "h3":
            elems.append(Paragraph(rich(b["text"]), ParagraphStyle("h3", parent=h2, fontSize=12)))
        elif b["type"] == "bullet":
            elems.append(Paragraph(rich(b["text"]), bullet_style, bulletText="•"))
        elif b["type"] == "code":
            elems.append(Paragraph(rich(b["text"]), ParagraphStyle("c", parent=body, fontName="Courier", fontSize=9)))
        elif b["type"] == "table" and b["rows"]:
            data = [[Paragraph(rich(c), body) for c in row] for row in b["rows"]]
            cols = max(len(r) for r in b["rows"])
            w = (A4[0] - 32 * mm) / cols
            t = Table(data, colWidths=[w] * cols)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F1FB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            elems.append(t)
            elems.append(Spacer(1, 2 * mm))
        else:
            elems.append(Paragraph(rich(b["text"]), body))

    doc.build(elems)
    return buf.getvalue()


def _md_to_docx(md: str, title: str) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor

    NAVY = RGBColor(0x1E, 0x27, 0x61)
    buf = io.BytesIO()
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)

    def add_runs(par, text: str):
        for frag, bold in _strip_bold(text):
            r = par.add_run(frag)
            r.bold = bold

    tp = doc.add_paragraph()
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(20)
    tr.font.color.rgb = NAVY
    meta = doc.add_paragraph()
    mr = meta.add_run(f"来源：知识库对话 · {title}")
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    for b in parse_md_blocks(md):
        t = b["type"]
        if t in ("h1", "h2", "h3"):
            level = {"h1": 1, "h2": 2, "h3": 3}[t]
            hp = doc.add_heading("", level=level)
            hp.style.font.color.rgb = NAVY
            add_runs(hp, b["text"])
        elif t == "bullet":
            bp = doc.add_paragraph(style="List Bullet")
            add_runs(bp, b["text"])
        elif t == "code":
            cp = doc.add_paragraph()
            cr = cp.add_run(b["text"])
            cr.font.name = "Consolas"
            cr.font.size = Pt(9)
        elif t == "table" and b["rows"]:
            cols = max(len(r) for r in b["rows"])
            table = doc.add_table(rows=len(b["rows"]), cols=cols)
            table.style = "Table Grid"
            for ri, row in enumerate(b["rows"]):
                for ci in range(cols):
                    cell_text = row[ci] if ci < len(row) else ""
                    pc = table.cell(ri, ci).paragraphs[0]
                    run = pc.add_run(cell_text.replace("**", ""))
                    run.font.size = Pt(9.5)
                    if ri == 0:
                        run.bold = True
        else:
            pp = doc.add_paragraph()
            add_runs(pp, b["text"])

    doc.save(buf)
    return buf.getvalue()


# ================= PPT：LLM 大纲 + python-pptx 渲染 =================

_OUTLINE_PROMPT = """你是资深 PPT 大纲设计师。请把下面的内容压缩为一份幻灯片大纲，要求：
1. 不超过 10 页内容页；每页标题不超过 15 字；每页 3~5 个要点，每条要点不超过 20 字；
2. 提炼结论与要点，不要照抄长句；信息过多时合并同类项；
3. 保留原文中的关键数字、术语与英文文案。

输出严格 JSON（不要 markdown 代码块、不要解释）：
{{"title":"演示标题","subtitle":"副标题","slides":[{{"title":"页标题","bullets":["要点1","要点2","要点3"],"note":"演讲备注一句话"}}]}}

内容：
{content}
"""


def _outline_via_llm(md: str) -> dict:
    from app.services.llm_service import llm_service

    prompt = _OUTLINE_PROMPT.replace("{content}", (md or "")[:12000])
    raw = llm_service.chat(prompt, temperature=0.2)
    m = re.search(r"\{.*\}", raw or "", re.S)
    if not m:
        raise ValueError("outline not json")
    data = json.loads(m.group(0))
    slides = [s for s in (data.get("slides") or []) if isinstance(s, dict) and s.get("title")]
    if not slides:
        raise ValueError("outline empty")
    return {"title": data.get("title") or "知识库总结", "subtitle": data.get("subtitle") or "", "slides": slides}


def _outline_fallback(md: str, title: str) -> dict:
    """LLM 大纲失败时的兜底：按二级标题切页"""
    slides, cur = [], None
    for b in parse_md_blocks(md):
        if b["type"] in ("h1", "h2"):
            if cur:
                slides.append(cur)
            cur = {"title": b["text"][:20], "bullets": [], "note": ""}
        elif b["type"] == "bullet" and cur is not None:
            cur["bullets"].append(b["text"][:30])
            if len(cur["bullets"]) >= 5:
                slides.append(cur)
                cur = {"title": cur["title"], "bullets": [], "note": ""}
    if cur and cur["bullets"]:
        slides.append(cur)
    return {"title": title, "subtitle": "", "slides": slides[:10] or [{"title": title, "bullets": ["（无内容）"], "note": ""}]}


def _md_to_pptx(md: str, title: str, use_llm: bool = True) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    outline = None
    if use_llm:
        try:
            outline = _outline_via_llm(md)
        except Exception as e:  # noqa: BLE001
            logger.warning("[knowledge_export] LLM 大纲失败，降级为标题切页: %s", e)
    if not outline:
        outline = _outline_fallback(md, title)

    NAVY = RGBColor(0x0F, 0x17, 0x2A)
    NAVY2 = RGBColor(0x1E, 0x27, 0x61)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    ACCENT = RGBColor(0x40, 0x9E, 0xFF)
    DARK = RGBColor(0x1E, 0x29, 0x3B)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    ICE = RGBColor(0xCA, 0xDC, 0xFC)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    font = "Microsoft YaHei"

    def add_rect(slide, x, y, w, h, color):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        return sh

    def add_text(slide, x, y, w, h, text, size, color, bold=False, center=False):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for li, ln in enumerate(str(text).split("\n")):
            para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            if center:
                para.alignment = 2
            run = para.add_run()
            run.text = ln
            run.font.name = font
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.bold = bold
        return tb

    # ---- 封面 ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, 0, Inches(3.4), W, Inches(0.05), ACCENT)
    add_text(s, Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.0), outline.get("title") or title, 40, WHITE, bold=True)
    add_text(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.6), outline.get("subtitle") or "知识库对话总结", 18, ICE)
    add_text(s, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4), "AI 驱动的研发质量前置平台 · 知识库对话导出", 12, MUTED)

    # ---- 内容页 ----
    for idx, sl in enumerate(outline["slides"], start=1):
        s = prs.slides.add_slide(blank)
        add_rect(s, 0, 0, W, Inches(0.9), NAVY2)
        add_text(s, Inches(0.5), Inches(0.14), Inches(11.0), Inches(0.6), sl.get("title") or f"第 {idx} 页", 24, WHITE, bold=True)
        add_text(s, Inches(12.2), Inches(0.2), Inches(0.9), Inches(0.5), f"{idx:02d}", 14, ICE)
        add_rect(s, 0, Inches(0.9), W, Inches(0.03), ACCENT)
        bullets = [str(b).strip() for b in (sl.get("bullets") or []) if str(b).strip()]
        body = "\n".join(f"•  {b}" for b in bullets) or "（本页无要点）"
        add_text(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(4.6), body, 18, DARK)
        note = (sl.get("note") or "").strip()
        if note:
            add_text(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5), f"备注：{note}", 11, MUTED)
        try:
            s.notes_slide.notes_text_frame.text = note
        except Exception:  # noqa: BLE001
            pass

    # ---- 尾页 ----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, W, H, NAVY)
    add_text(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.0), "谢谢", 44, WHITE, bold=True)
    add_text(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.5), "由 AI 驱动的研发质量前置平台 · 知识库对话自动生成", 14, ICE)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_chat(answer_md: str, title: str, fmt: str, use_llm: bool = True) -> tuple[bytes, str, str]:
    """统一出口：返回 (字节流, MIME, 文件名)"""
    safe = re.sub(r"[\\/:*?\"<>|\s]+", "_", (title or "知识库总结").strip())[:60] or "知识库总结"
    if fmt == "word":
        return _md_to_docx(answer_md, title), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{safe}.docx"
    if fmt == "ppt":
        return _md_to_pptx(answer_md, title, use_llm), "application/vnd.openxmlformats-officedocument.presentationml.presentation", f"{safe}.pptx"
    if fmt == "pdf":
        return _md_to_pdf(answer_md, title), "application/pdf", f"{safe}.pdf"
    raise ValueError(f"不支持的导出格式: {fmt}")
